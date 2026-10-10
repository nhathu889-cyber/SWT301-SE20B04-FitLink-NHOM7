import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent / '.slide-tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, GRAY, BG = '172B4D', '52627A', 'F5F7FB'

def box(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape

def text(slide, x, y, w, h, value, size=20, color=NAVY, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = 'Arial'; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(9)
    return shape

def base(title, subtitle, count, color, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    box(slide, .48, .48, .09, 1.05, color)
    text(slide, .78, .38, 10.5, .65, title, 32, bold=True)
    text(slide, .8, 1.12, 10.7, .5, subtitle, 17, GRAY)
    box(slide, 11.3, .45, 1.48, .8, color)
    text(slide, 11.45, .58, 1.2, .5, count, 25, 'FFFFFF', True)
    text(slide, .65, 6.86, 11.6, .45, 'FITLINK  |  Round 1 · 07/10/2026  |  Jest unit test · Mock dữ liệu, không ghi DB thật', 12, GRAY)
    text(slide, 12.35, 6.87, .4, .3, str(number), 13, GRAY)
    return slide

s = base('CASE ĐẠT', '9/24 case đạt trong phạm vi kiểm thử đã thực hiện.', '9 ĐẠT', '15805D', 1)
cards = [
    (.65, 5.8, 'AUTHENTICATION · 6 CASE', 'TC_AUTH_002 — Chặn dữ liệu đăng ký sai*\nTC_AUTH_003 — Từ chối email đăng ký trùng\nTC_AUTH_004 — Đăng nhập đúng, cấp cookie\nTC_AUTH_005 — Chặn mật khẩu / đầu vào sai\nTC_AUTH_007 — Chặn thiếu phiên / sai quyền\nTC_AUTH_008 — Chặn token sai / hết hạn'),
    (6.7, 5.98, 'BOOKING & SCHEDULE · 3 CASE', 'TC_BOOK_001 — Tạo booking đúng dữ liệu\nTC_BOOK_002 — Không lưu booking thiếu / sai dữ liệu\nTC_SCH_001 — Sinh lịch tương lai hợp lệ')]
for x,w,title,body in cards:
    box(s,x,1.95,w,3.95,'FFFFFF'); text(s,x+.22,2.15,w-.44,.45,title,18,'15805D',True)
    text(s,x+.22,2.8,w-.44,2.95,body,18)
text(s,.8,6.08,11.8,.57,'* AUTH_002: kiểm tra middleware riêng; chưa kiểm tra wiring route đăng ký.\nAUTH_004–005: điều chỉnh theo đăng nhập bằng số điện thoại + cookie HttpOnly.',13,GRAY)

s = base('CASE KHÔNG ĐẠT', '5 case phát hiện hành vi chưa đáp ứng kỳ vọng trong testcase.', '5 LỖI', 'C33F49', 2)
rows = [
    ('TC_BOOK_003', 'Đặt trùng slot', 'Booking thứ hai vẫn được chấp nhận (HTTP 200).'),
    ('TC_SCH_002', 'Ngày bắt đầu sai', 'Vòng lặp không dừng; tiến trình thử bị ngắt sau 1.500 ms.'),
    ('TC_SCH_003', 'Giờ kết thúc ≤ bắt đầu', 'Không từ chối slot.end sai; dùng duration để tính giờ kết thúc.'),
    ('TC_SCH_004', 'Lịch trong quá khứ', 'Vẫn tạo lịch (HTTP 201) khi carryForward=false.'),
    ('TC_SCH_005', 'Lịch bị chồng lấn', 'Sinh lịch trùng lần hai vẫn thành công (HTTP 201).')]
for i,(ident,title,detail) in enumerate(rows):
    y=1.85+i*.88
    box(s,.65,y,12.03,.77,'FFFFFF')
    text(s,.87,y+.12,2.1,.45,ident,18,'C33F49',True)
    text(s,3.02,y+.12,3.08,.45,title,18,bold=True)
    text(s,6.17,y+.12,6.15,.53,detail,17)
text(s,.8,6.35,11.8,.35,'Kết quả của round kiểm thử; các lỗi này chưa được sửa trong yêu cầu lập báo cáo.',13,GRAY)

s = base('CASE CHƯA ĐỦ ĐIỀU KIỆN', '10 case Pending — chưa thực thi đầy đủ; không tính là đạt hoặc không đạt.', '10 CHỜ', 'B97812', 3)
groups = [
    (.65,1.85,5.8,'TC_AUTH_001 · 1 CASE','Đăng ký tài khoản hoàn chỉnh','Chưa kiểm thử toàn bộ vòng đời xác nhận email và tạo tài khoản.'),
    (6.7,1.85,5.98,'TC_AUTH_006 · 1 CASE','Thông tin tài khoản','Đã có API test riêng, nhưng chưa ánh xạ đầy đủ vào case của round này.'),
    (.65,4.05,5.8,'TC_BOOK_004–008 · 5 CASE','Xem / sửa / hủy booking & quyền sở hữu','Chưa có API đọc, sửa, hủy booking tương ứng để chạy đúng quy trình testcase.'),
    (6.7,4.05,5.98,'TC_SCH_006–008 · 3 CASE','Sửa thời gian lịch & quyền sở hữu','Chưa có hợp đồng API đọc/sửa thời gian tương ứng. API đổi trạng thái buổi tập là thao tác khác.')]
for x,y,w,ident,title,detail in groups:
    box(s,x,y,w,1.93,'FFFFFF')
    text(s,x+.22,y+.15,w-.44,.35,ident,17,'B97812',True)
    text(s,x+.22,y+.62,w-.44,.4,title,18,bold=True)
    text(s,x+.22,y+1.08,w-.44,.7,detail,16,GRAY)
text(s,.8,6.35,11.8,.35,'Cần bổ sung API / điều kiện xác nhận và ánh xạ kiểm thử trước khi chuyển Pending thành kết quả thực thi.',13,GRAY)

output = ROOT / 'FitLink_Test_Cases_3_Slides.pptx'
prs.save(output)
assert len(Presentation(output).slides) == 3
print(output)
